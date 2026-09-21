"""Narrow workspace helper executed only inside a disposable OCI child container."""

from __future__ import annotations

import base64
import codecs
import json
import os
import subprocess
import sys
import tempfile
import hashlib
import shutil
import re
import secrets
import stat
from pathlib import Path
from typing import Any

ROOT = Path("/workspace")
SOURCE = Path("/source")
MAX_INPUT_BYTES = 500_000
MAX_GIT_BYTES = 2_000_000
MAX_PREPARED_BUNDLE_BYTES = 16 * 1024 * 1024
MAX_PREPARED_BUNDLE_CHUNK_BYTES = 700_000


def _fail(code: str) -> None:
    print(json.dumps({"error": code}, separators=(",", ":")), file=sys.stderr)
    raise SystemExit(2)


def _path(value: str, allow_root: bool = False) -> Path:
    if not value or value.startswith("/") or "\x00" in value:
        _fail("WORKSPACE_PATH_INVALID")
    relative = Path(value)
    if any(part.casefold() == ".git" for part in relative.parts):
        _fail("WORKSPACE_GIT_METADATA_FORBIDDEN")
    if any(part in ("", ".", "..") for part in relative.parts):
        if not (allow_root and value == "."):
            _fail("WORKSPACE_PATH_INVALID")
    root = ROOT.resolve(strict=True)
    candidate = root if value == "." else root.joinpath(relative)
    current = root
    for part in (() if value == "." else relative.parts):
        current = current / part
        if current.is_symlink():
            _fail("WORKSPACE_SYMLINK_FORBIDDEN")
    resolved = candidate.resolve(strict=False)
    if resolved != root and root not in resolved.parents:
        _fail("WORKSPACE_PATH_INVALID")
    return candidate


def _input() -> bytes:
    encoded = os.environ.pop("SPACEAGENT_INPUT_BASE64", "")
    try:
        value = base64.b64decode(encoded, validate=True)
    except Exception:
        _fail("WORKSPACE_INPUT_INVALID")
    if len(value) > MAX_INPUT_BYTES:
        _fail("WORKSPACE_INPUT_TOO_LARGE")
    return value


def _git(*arguments: str, limit: int = MAX_GIT_BYTES, accepted: tuple[int, ...] = (0,)) -> str:
    try:
        result = subprocess.run(
            ["git", *arguments], cwd=ROOT, stdin=subprocess.DEVNULL,
            stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=60, check=False,
            env={"PATH": os.environ.get("PATH", "/usr/bin:/bin"),
                 "HOME": "/home/sandbox", "GIT_TERMINAL_PROMPT": "0",
                 "GIT_OPTIONAL_LOCKS": "0"},
        )
    except Exception:
        _fail("WORKSPACE_GIT_FAILED")
    if result.returncode not in accepted:
        _fail("WORKSPACE_GIT_FAILED")
    if len(result.stdout) + len(result.stderr) > limit:
        _fail("WORKSPACE_GIT_OUTPUT_LIMIT")
    return result.stdout.decode("utf-8", "replace")


def _changed() -> list[str]:
    status = _git("status", "--porcelain", "--untracked-files=all", limit=200_000)
    return sorted({line[3:] for line in status.splitlines() if len(line) > 3})


def _read(path: str, maximum: int) -> dict[str, Any]:
    target = _path(path)
    if not target.is_file():
        _fail("WORKSPACE_FILE_NOT_FOUND")
    limit = max(1, min(maximum, 200_000))
    size = target.stat().st_size
    # A character cap is a transport AND allocation bound, not a post-read slice.
    with target.open("rb") as stream:
        raw = stream.read(limit * 4 + 4)
    decoder = codecs.getincrementaldecoder("utf-8")("replace")
    content = decoder.decode(raw, final=len(raw) >= size)
    return {"path": path, "content": content[:limit], "sizeBytes": size,
            "truncated": size > len(raw) or len(content) > limit}


def _list(path: str, depth: int, maximum: int) -> dict[str, Any]:
    base_value = path or "."
    base = _path(base_value, allow_root=True)
    if not base.is_dir():
        _fail("WORKSPACE_DIRECTORY_NOT_FOUND")
    root = ROOT.resolve(strict=True)
    entries: list[dict[str, Any]] = []
    truncated = False
    for current, directories, files in os.walk(base, followlinks=False):
        current_path = Path(current)
        relative_depth = len(current_path.relative_to(base).parts)
        directories[:] = sorted(name for name in directories
                                if name.casefold() != ".git" and not (current_path / name).is_symlink())
        if relative_depth >= max(0, min(depth, 8)):
            directories[:] = []
        for name, directory in [(name, True) for name in directories] + [
                (name, False) for name in sorted(files) if name.casefold() != ".git"]:
            target = current_path / name
            if target.is_symlink():
                continue
            relative = target.relative_to(root).as_posix()
            size = 0 if directory else target.stat().st_size
            entries.append({"path": relative, "directory": directory, "sizeBytes": size})
            if len(entries) >= max(1, min(maximum, 500)):
                truncated = True
                return {"root": base_value, "entries": entries, "truncated": truncated}
    return {"root": base_value, "entries": entries, "truncated": truncated}


def _snapshot(maximum: int) -> dict[str, Any]:
    limit = max(1_000, min(maximum, 2_000_000))
    head = _git("rev-parse", "HEAD", limit=1_000).strip().lower()
    status = _git("status", "--porcelain", "--untracked-files=all", limit=200_000)
    patch = _git("diff", "--binary", "--no-ext-diff", limit=limit)
    untracked = [line for line in status.splitlines() if line.startswith("?? ")]
    if len(untracked) > 100:
        _fail("WORKSPACE_UNTRACKED_FILE_LIMIT")
    for line in untracked:
        value = _git("diff", "--binary", "--no-index", "/dev/null", line[3:],
                     limit=limit, accepted=(0, 1))
        if len((patch + value).encode("utf-8")) > limit:
            return {"headCommit": head, "status": status, "patch": patch,
                    "changedFiles": _changed(), "truncated": True}
        patch += value
    return {"headCommit": head, "status": status, "patch": patch,
            "changedFiles": _changed(), "truncated": False}


def _write(path: str, value: bytes) -> dict[str, Any]:
    target = _path(path)
    parent = target.parent
    parent.mkdir(parents=True, exist_ok=True)
    _path(parent.resolve().relative_to(ROOT.resolve()).as_posix() or ".", allow_root=True)
    descriptor, temporary = tempfile.mkstemp(prefix=".spaceagent-", dir=parent)
    try:
        with os.fdopen(descriptor, "wb") as stream:
            stream.write(value)
        os.replace(temporary, target)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    return {"path": path, "sizeBytes": len(value), "changedFiles": _changed()}


def _document_workspace_write(path: str, value: bytes) -> dict[str, Any]:
    target = _path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    _path(target.parent.resolve().relative_to(ROOT.resolve()).as_posix() or ".", allow_root=True)
    descriptor, temporary = tempfile.mkstemp(prefix=".spaceagent-", dir=target.parent)
    try:
        with os.fdopen(descriptor, "wb") as stream:
            stream.write(value)
        os.replace(temporary, target)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    return {"path": path, "sizeBytes": len(value),
            "contentHash": "sha256:" + hashlib.sha256(value).hexdigest()}


def _document_workspace_delete(path: str) -> dict[str, Any]:
    target = _path(path)
    if target.exists() and not target.is_file():
        _fail("WORKSPACE_DELETE_INVALID")
    target.unlink(missing_ok=True)
    return {"path": path, "sizeBytes": 0,
            "contentHash": "sha256:" + hashlib.sha256(b"").hexdigest()}


def _document_workspace_proof(path: str) -> dict[str, Any]:
    target = _path(path)
    if not target.exists():
        return {"path": path, "exists": False, "sizeBytes": 0,
                "contentHash": "sha256:" + hashlib.sha256(b"").hexdigest()}
    if not target.is_file():
        _fail("WORKSPACE_FILE_NOT_FOUND")
    value = target.read_bytes()
    return {"path": path, "exists": True, "sizeBytes": len(value),
            "contentHash": "sha256:" + hashlib.sha256(value).hexdigest()}


def _document_workspace_clear() -> dict[str, Any]:
    root = ROOT.resolve(strict=True)
    count = 0
    for target in sorted(root.rglob("*"), key=lambda value: len(value.parts), reverse=True):
        if target.is_symlink():
            _fail("WORKSPACE_SYMLINK_FORBIDDEN")
        if target.is_file():
            target.unlink()
            count += 1
        elif target.is_dir():
            target.rmdir()
    return {"deletedFiles": count}


def _document_read(path: str, maximum: int) -> dict[str, Any]:
    target = _path(path)
    suffix = target.suffix.lower()
    if suffix == ".docx":
        from docx import Document
        content = "\n".join(paragraph.text for paragraph in Document(target).paragraphs)
        media = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif suffix == ".pdf":
        from pypdf import PdfReader
        content = "\n".join(page.extract_text() or "" for page in PdfReader(target).pages[:200])
        media = "application/pdf"
    elif suffix in (".html", ".htm"):
        from bs4 import BeautifulSoup
        content = BeautifulSoup(target.read_text("utf-8", errors="replace"), "html.parser").get_text("\n")
        media = "text/html"
    else:
        content = target.read_text("utf-8", errors="replace")
        media = "text/markdown" if suffix in (".md", ".markdown") else "text/plain"
    limit = max(1, min(maximum, 200_000))
    return {"path": path, "mediaType": media, "content": content[:limit],
            "metadata": {}, "truncated": len(content) > limit}


def _document_write(path: str, format_name: str) -> dict[str, Any]:
    content = _input().decode("utf-8", "strict")
    if format_name == "docx":
        from docx import Document
        document = Document()
        for line in content.splitlines() or [""]:
            document.add_paragraph(line)
        target = _path(path)
        target.parent.mkdir(parents=True, exist_ok=True)
        _path(target.parent.resolve().relative_to(ROOT.resolve()).as_posix() or ".", allow_root=True)
        document.save(target)
        result = {"path": path, "sizeBytes": target.stat().st_size,
                  "changedFiles": _changed()}
    elif format_name in ("text", "markdown", "html"):
        result = _write(path, content.encode("utf-8"))
    else:
        _fail("WORKSPACE_DOCUMENT_FORMAT_INVALID")
    result["format"] = format_name
    return result


def _prepare_commit(expected_base: str, patch_hash: str) -> dict[str, Any]:
    if re.fullmatch(r"[0-9a-f]{40,64}", expected_base or "") is None:
        _fail("MERGE_BASE_INVALID")
    if re.fullmatch(r"sha256:[0-9a-f]{64}", patch_hash or "") is None:
        _fail("MERGE_PATCH_HASH_INVALID")
    message = _input().decode("utf-8", "strict").strip()
    if not message or len(message) > 500:
        _fail("WORKSPACE_COMMIT_MESSAGE_INVALID")
    snapshot = _snapshot(MAX_GIT_BYTES)
    if snapshot["headCommit"] != expected_base:
        parent = _git("rev-parse", snapshot["headCommit"] + "^", limit=1_000).strip().lower()
        body = _git("log", "-1", "--format=%B", snapshot["headCommit"], limit=8_000)
        if (parent == expected_base and f"SpaceAgent-Patch-Hash: {patch_hash}" in body
                and not snapshot["status"].strip()):
            return _prepared_bundle(snapshot["headCommit"], expected_base)
        _fail("MERGE_WORKSPACE_HEAD_DRIFT")
    if snapshot["truncated"]:
        _fail("MERGE_PROPOSAL_BASE_MISMATCH")
    actual_hash = "sha256:" + hashlib.sha256(snapshot["patch"].encode("utf-8")).hexdigest()
    if actual_hash != patch_hash:
        _fail("MERGE_PATCH_HASH_MISMATCH")
    _git("add", "-A", limit=200_000)
    _git("-c", "user.name=SpaceAgent", "-c", "user.email=spaceagent@localhost",
         "commit", "--no-gpg-sign", "-m", message, "-m",
         f"SpaceAgent-Patch-Hash: {patch_hash}", limit=200_000)
    commit = _git("rev-parse", "HEAD", limit=1_000).strip().lower()
    return _prepared_bundle(commit, expected_base)


def _prepared_bundle(commit: str, expected_base: str) -> dict[str, Any]:
    if re.fullmatch(r"[0-9a-f]{40,64}", commit or "") is None:
        _fail("MERGE_COMMIT_INVALID")
    git_dir = ROOT / ".git"
    if git_dir.is_symlink() or not git_dir.is_dir():
        _fail("MERGE_GIT_DIRECTORY_INVALID")
    transfer_dir = git_dir / "spaceagent-transfer"
    if transfer_dir.exists() and (transfer_dir.is_symlink() or not transfer_dir.is_dir()):
        _fail("MERGE_BUNDLE_PATH_INVALID")
    transfer_dir.mkdir(mode=0o700, exist_ok=True)
    if transfer_dir.is_symlink():
        _fail("MERGE_BUNDLE_PATH_INVALID")
    for previous in transfer_dir.iterdir():
        if previous.is_symlink() or not previous.is_file():
            _fail("MERGE_BUNDLE_PATH_INVALID")
        previous.unlink()
    bundle = transfer_dir / f"{commit}.bundle"
    temporary = transfer_dir / f".{commit}.{secrets.token_hex(16)}.partial"
    try:
        _git("bundle", "create", str(temporary), "HEAD", f"^{expected_base}",
             limit=200_000)
        metadata = temporary.lstat()
        if not stat.S_ISREG(metadata.st_mode) or metadata.st_size < 1 \
                or metadata.st_size > MAX_PREPARED_BUNDLE_BYTES:
            _fail("MERGE_BUNDLE_SIZE_INVALID")
        os.chmod(temporary, 0o600)
        os.replace(temporary, bundle)
        metadata = bundle.lstat()
        if not stat.S_ISREG(metadata.st_mode) or metadata.st_size < 1 \
                or metadata.st_size > MAX_PREPARED_BUNDLE_BYTES:
            _fail("MERGE_BUNDLE_SIZE_INVALID")
        digest = hashlib.sha256()
        with bundle.open("rb") as stream:
            for block in iter(lambda: stream.read(64 * 1024), b""):
                digest.update(block)
        return {
            "commit": commit,
            "bundleReference": f".git/spaceagent-transfer/{commit}.bundle",
            "bundleSha256": "sha256:" + digest.hexdigest(),
            "bundleBytes": metadata.st_size,
        }
    finally:
        temporary.unlink(missing_ok=True)


def _read_prepared_bundle(commit: str, offset_value: str, maximum_value: str) -> dict[str, Any]:
    if re.fullmatch(r"[0-9a-f]{40,64}", commit or "") is None:
        _fail("MERGE_COMMIT_INVALID")
    try:
        offset = int(offset_value)
        maximum = int(maximum_value)
    except ValueError:
        _fail("MERGE_BUNDLE_READ_INVALID")
    git_dir = ROOT / ".git"
    transfer_dir = git_dir / "spaceagent-transfer"
    bundle = transfer_dir / f"{commit}.bundle"
    if git_dir.is_symlink() or not git_dir.is_dir() \
            or transfer_dir.is_symlink() or not transfer_dir.is_dir() \
            or bundle.is_symlink():
        _fail("MERGE_BUNDLE_PATH_INVALID")
    metadata = bundle.lstat()
    if not stat.S_ISREG(metadata.st_mode) or metadata.st_size < 1 \
            or metadata.st_size > MAX_PREPARED_BUNDLE_BYTES \
            or offset < 0 or offset >= metadata.st_size \
            or maximum < 1 or maximum > MAX_PREPARED_BUNDLE_CHUNK_BYTES:
        _fail("MERGE_BUNDLE_READ_INVALID")
    expected = min(maximum, metadata.st_size - offset)
    with bundle.open("rb") as stream:
        stream.seek(offset)
        chunk = stream.read(maximum)
    if len(chunk) != expected:
        _fail("MERGE_BUNDLE_READ_INVALID")
    return {
        "commit": commit,
        "bundleReference": f".git/spaceagent-transfer/{commit}.bundle",
        "offset": offset,
        "totalBytes": metadata.st_size,
        "chunkBase64": base64.b64encode(chunk).decode("ascii"),
    }


def _materialize_snapshot() -> dict[str, Any]:
    source = SOURCE.resolve(strict=True)
    root = ROOT.resolve(strict=True)
    if any(root.iterdir()):
        _fail("WORKSPACE_NOT_EMPTY")
    count = 0
    for current, directories, files in os.walk(source, followlinks=False):
        current_path = Path(current)
        if any((current_path / name).is_symlink() for name in directories + files):
            _fail("SOURCE_SNAPSHOT_SYMLINK_FORBIDDEN")
        target = root / current_path.relative_to(source)
        target.mkdir(parents=True, exist_ok=True)
        for name in sorted(files):
            shutil.copyfile(current_path / name, target / name)
            count += 1
    _git("init", "-b", "main", limit=20_000)
    _git("add", "-A", limit=20_000)
    _git("-c", "user.name=SpaceAgent", "-c", "user.email=spaceagent@localhost",
         "commit", "--no-gpg-sign", "-m", "Materialized managed snapshot", limit=20_000)
    return {"headCommit": _git("rev-parse", "HEAD", limit=1_000).strip().lower(),
            "fileCount": count}


def main() -> None:
    if len(sys.argv) < 2:
        _fail("WORKSPACE_OPERATION_REQUIRED")
    operation = sys.argv[1]
    if operation == "file-read" and len(sys.argv) == 4:
        result = _read(sys.argv[2], int(sys.argv[3]))
    elif operation == "file-list" and len(sys.argv) == 5:
        result = _list(sys.argv[2], int(sys.argv[3]), int(sys.argv[4]))
    elif operation == "git" and len(sys.argv) == 3:
        result = _snapshot(int(sys.argv[2]))
    elif operation == "write-file" and len(sys.argv) == 3:
        result = _write(sys.argv[2], _input())
    elif operation == "delete-file" and len(sys.argv) == 3:
        target = _path(sys.argv[2])
        if target.exists() and not target.is_file():
            _fail("WORKSPACE_DELETE_INVALID")
        target.unlink(missing_ok=True)
        result = {"path": sys.argv[2], "sizeBytes": 0, "changedFiles": _changed()}
    elif operation == "document-file-write" and len(sys.argv) == 3:
        result = _document_workspace_write(sys.argv[2], _input())
    elif operation == "document-file-delete" and len(sys.argv) == 3:
        result = _document_workspace_delete(sys.argv[2])
    elif operation == "document-file-proof" and len(sys.argv) == 3:
        result = _document_workspace_proof(sys.argv[2])
    elif operation == "document-file-clear" and len(sys.argv) == 2:
        result = _document_workspace_clear()
    elif operation == "document-read" and len(sys.argv) == 4:
        result = _document_read(sys.argv[2], int(sys.argv[3]))
    elif operation == "document-write" and len(sys.argv) == 4:
        result = _document_write(sys.argv[2], sys.argv[3])
    elif operation == "prepare-commit" and len(sys.argv) == 4:
        result = _prepare_commit(sys.argv[2], sys.argv[3])
    elif operation == "prepared-bundle-read" and len(sys.argv) == 5:
        result = _read_prepared_bundle(sys.argv[2], sys.argv[3], sys.argv[4])
    elif operation == "snapshot-materialize" and len(sys.argv) == 2:
        result = _materialize_snapshot()
    else:
        _fail("WORKSPACE_OPERATION_INVALID")
    print(json.dumps(result, ensure_ascii=False, separators=(",", ":")))


if __name__ == "__main__":
    main()
